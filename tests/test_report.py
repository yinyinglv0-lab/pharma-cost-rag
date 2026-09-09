# -*- coding: utf-8 -*-
"""模块一验收：模板解析注册表全覆盖 + 渲染零残留 + 数字与金丝雀一致 + 图表嵌入。"""
import io
import re

import docx
import pytest

from app.agents.graph import run_agent
from app.agents.llm import MockLLM
from app.report.registry import parse_template
from app.report.render import build_report

PH = re.compile(r"\{\{.*?\}\}")


def test_registry_covers_template():
    reg = parse_template()
    assert len(reg) == 101, f"占位符数 {len(reg)}"
    assert reg["产品名称"] == "数值" and reg["材料成本归因分析文本"] == "文本"
    assert reg["原材料成本明细表格"] == "表格"
    assert reg["差异结构拆解分析"] == "文本", "拆解分析应为文本类"


def test_report_no_residue_and_numbers_match():
    out = run_agent("请生成银黄口服液2026-05的月度成本分析报告",
                    session_id="rep-test", llm=MockLLM())
    data = build_report("银黄口服液", "2026-05", draft=out["draft"],
                        rpa_task=out["rpa_task"], llm=MockLLM(), charts=True)
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    assert not PH.search(text), f"残留占位符: {PH.findall(text)[:5]}"
    # 金丝雀数字：银黄 2026-05 单位成本 11.21 / 环比 +2.84%
    assert "11.21" in text, "本月单位成本未填入"
    assert "银黄口服液" in text and "2026-05" in text
    # 表格内容
    assert "金银花" in text, "原材料明细表格未渲染"
    # 图表嵌入（趋势/瀑布/结构 三张 PNG）
    assert len(d.inline_shapes) >= 3, f"嵌入图片数 {len(d.inline_shapes)}"


def test_report_all_scenarios_no_residue():
    for product, month in [("银黄口服液", "2026-05"), ("板蓝根颗粒", "2026-02"),
                           ("六味地黄胶囊", "2026-03")]:
        out = run_agent(f"请生成{product}{month}的月度成本分析报告",
                        session_id=f"rep-{product}-{month}", llm=MockLLM())
        data = build_report(product, month, draft=out["draft"], rpa_task=out["rpa_task"],
                            llm=MockLLM(), charts=False)
        d = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in d.paragraphs)
        assert not PH.search(text), f"{product} 残留占位符"
