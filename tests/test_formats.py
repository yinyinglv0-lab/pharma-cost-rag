# -*- coding: utf-8 -*-
"""Word/TXT 格式支持验收（赛题 5.1.2：知识库文档格式支持 PDF、Word、TXT）。"""
from app.rag.parse import build_office_chunks
from app.rag.store import HybridRetriever

DOCX_TEXT = """一、产品基本信息
本测试文档用于验证 Word 格式知识文档支持。

2.1 处方量
金银花 5.00kg 黄芩提取物 2.00kg

2.3 单盒材料成本参考
金银花 128 元/kg 黄芩提取物 45 元/kg

三、功能主治
清热解毒，消炎。
"""


def _write_docx(path):
    import docx
    d = docx.Document()
    for para in DOCX_TEXT.splitlines():
        if para.strip():
            d.add_paragraph(para)
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "原料"
    t.rows[0].cells[1].text = "处方量(kg)"
    t.rows[1].cells[0].text = "金银花"
    t.rows[1].cells[1].text = "5.00"
    d.save(str(path))


def test_docx_parsing(tmp_path):
    _write_docx(tmp_path / "测试工艺.docx")
    chunks = build_office_chunks(tmp_path)
    assert chunks, "docx 未解析出知识块"
    types = {c["type"] for c in chunks}
    assert "处方量" in types and "处方参考价" in types
    for c in chunks:
        assert c["source"].endswith(".docx") and c["chunk_id"] and c["text"]
    assert any("5.00" in c["text"] for c in chunks), "表格文本应进入知识块"


def test_txt_parsing(tmp_path):
    (tmp_path / "GMP补充说明.txt").write_text(DOCX_TEXT, encoding="utf-8")
    chunks = build_office_chunks(tmp_path)
    assert chunks and any(c["source"].endswith(".txt") for c in chunks)
    assert any(c["type"] == "处方量" for c in chunks)


def test_office_chunks_retrievable(tmp_path):
    """docx 知识块进三通道检索，结果可溯源（来源标注条款覆盖 Word 格式）。"""
    _write_docx(tmp_path / "银黄补充配方.docx")
    chunks = build_office_chunks(tmp_path)
    kb = HybridRetriever(chunks, embedder="hash")
    results, _ = kb.retrieve("银黄口服液处方量", top_k=3)
    assert results, "docx 块未被检索到"
    assert any(c["type"] == "处方量" for c in results)
    for c in results:
        assert c["source"] and c["chunk_id"], "来源标注缺失"
