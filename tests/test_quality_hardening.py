# -*- coding: utf-8 -*-
"""交付级质量加固测试：四维度（溯源/图表标注/闭环回写/对标RPA计划）+ 边界探针。"""
import io

import pytest

import docx

from app.agents.graph import run_agent
from app.agents.guard import check_rpa_task
from app.agents.llm import MockLLM
from app.benchmark.steps import step3_cause
from app.core.calc import CalculationService
from app.core.datastore import DataStore
from app.dashboard.charts import CHART_BUILDERS
from app.report.render import build_report
from app.rpa.closure import closure_records, record_closure


# ============ 维度1: 数据 100% 正确 ============
def test_provenance_row_numbers_and_formulas():
    svc = CalculationService(DataStore())
    prov = svc.provenance("银黄口服液", "2026-05")
    assert len(prov) >= 5
    for p in prov:
        assert p["metric"] and p["formula"] and p["csv_file"]
    uc = next(p for p in prov if p["metric"] == "单位成本")
    assert uc["value"] == 11.21, "金丝雀数字误差必须为 0"
    assert uc["line_number"] > 1, "必须给出 CSV 行号"


def test_report_contains_provenance_appendix():
    out = run_agent("请生成银黄口服液2026-05的月度成本分析报告",
                    session_id="qa-prov", llm=MockLLM())
    data = build_report("银黄口服液", "2026-05", draft=out["draft"],
                        rpa_task=out["rpa_task"], llm=MockLLM(), charts=False)
    d = docx.Document(io.BytesIO(data))
    text = "\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                text += "\n" + c.text
    assert "数字溯源表" in text and "成本汇总_2026年1-6月.csv" in text
    assert "11.21" in text


def test_identities_in_report_values():
    svc = CalculationService(DataStore())
    for p in ("银黄口服液", "板蓝根颗粒", "六味地黄胶囊"):
        for m in range(2, 7):
            mm = f"2026-{m:02d}"
            tf = svc.three_factor(p, mm)
            assert abs(sum(tf[e]["share_pct"] for e in ("材料", "人工", "制费")) - 100) < 1e-9, \
                f"{p} {mm} 贡献度恒等式破坏"
            assert abs(sum(tf[e]["delta"] for e in ("材料", "人工", "制费"))
                       - tf["unit_cost_delta"]) < 1e-9


# ============ 维度2: 图表科学严谨 ============
@pytest.mark.parametrize("builder", ["趋势", "瀑布", "结构", "热力图"])
def test_chart_annotations(builder):
    if builder == "趋势":
        opt = CHART_BUILDERS[builder]("银黄口服液")
    elif builder == "瀑布":
        opt = CHART_BUILDERS[builder]("银黄口服液", "2026-02")
    elif builder == "结构":
        opt = CHART_BUILDERS[builder]("银黄口服液", "2026-05")
    else:
        opt = CHART_BUILDERS[builder]()
    sub = opt["title"].get("subtext", "")
    assert "数据来源" in sub and "元/盒" in sub and "2026-01~2026-06" in sub, \
        f"{builder} 缺单位/来源/时间范围标注"


# ============ 维度3: RPA 闭环 ============
def test_rpa_schema_full_alignment():
    good = {
        "task_id": "TASK-202605-0001", "task_title": "核查金银花采购合同调价条款",
        "assignee": {"name": "张伟", "department": "采购部", "role": "采购经理"},
        "source": {"analysis_type": "月度成本分析", "analysis_month": "2026-05",
                   "product": "银黄口服液", "finding": "金银花涨价"},
        "priority": "high", "deadline": "2026-06-03",
        "suggestion": "核查调价条款", "notify_method": "wechat",
        "created_at": "2026-06-01T09:00:00",
    }
    assert check_rpa_task(good).passed
    bad = dict(good)
    bad["source"] = {"analysis_month": "2026-05", "product": "银黄口服液"}  # 缺 analysis_type/finding
    g = check_rpa_task(bad)
    assert not g.passed and any("analysis_type" in v for v in g.violations)


def test_closure_write_back(tmp_path, monkeypatch):
    monkeypatch.setenv("AGENT_ARTIFACT_DIR", str(tmp_path))
    task = {
        "task_id": "TASK-202605-0001-DONE", "task_title": "核查金银花采购合同调价条款",
        "assignee": {"name": "张伟", "department": "采购部"},
        "source": {"analysis_type": "月度成本分析", "analysis_month": "2026-05",
                   "product": "银黄口服液", "finding": "金银花涨价"},
        "priority": "high", "deadline": "2026-06-03",
        "suggestion": "核查调价条款", "notify_method": "wechat",
        "created_at": "2026-06-01T09:00:00",
    }
    rec = record_closure(task)
    assert rec["task_id"].endswith("-DONE") and rec["status"] == "completed"
    assert abs(rec["net_effect_pct"] - (-0.37)) < 0.05, "闭环净效应应与 DiD 一致"
    assert len(rec["warnings"]) == 2
    assert any(r["task_id"] == task["task_id"] for r in closure_records())


# ============ 维度4: 对标审计 ============
def test_rpa_plan_with_assignee_priority_deadline():
    c = step3_cause("银黄口服液")
    plan = c["rpa_plan"]
    assert plan, "建议必须可转 RPA 计划"
    for item in plan:
        for field in ("suggestion", "assignee", "priority", "deadline"):
            assert field in item and item[field], f"RPA 计划缺字段 {field}"
    assert plan[0]["priority"] in ("high", "medium", "low")


def test_dual_caliber_explicit():
    svc = CalculationService(DataStore())
    u, t = svc.yoy_unit_cost_overall(), svc.yoy_total_cost_overall()
    assert abs(u - 2.43) < 0.03 and abs(t - 7.77) < 0.03
    assert abs(u - t) > 1.0, "两口径必须显式区分（不得混用）"


# ============ 边界探针 ============
def test_boundary_prev_month_guard():
    svc = CalculationService(DataStore())
    # 无上月 → 标记 null（None），而非 0 或静默错数
    assert svc.three_factor("银黄口服液", "2026-01") is None
    assert svc.mom("银黄口服液", "2026-01", "材料") is None


def test_boundary_zero_price_pct_no_noise_bar():
    from app.simulator.engine import simulate_price
    sc = simulate_price("银黄口服液", "金银花", 0.0, "2026-05")
    assert sc.factors == [] and sc.total_delta == 0.0, "零增量不得产生因素条"
